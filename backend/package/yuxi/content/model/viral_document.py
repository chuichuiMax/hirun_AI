"""以解析原文单元作为边界：Agent 选择单元，代码提取原文，不接受模型重写正文。"""

import csv
import html
import io
import re

from pydantic import Field, StrictInt

from yuxi.content.model.viral_assets import ViralContract


class DetectedArticle(ViralContract):
    title_unit: StrictInt
    body_units: list[StrictInt] = Field(min_length=1)
    industry_slug: str = Field(min_length=1)
    reason: str = Field(min_length=1)


class ViralDocumentResultV1(ViralContract):
    articles: list[DetectedArticle] = Field(default_factory=list, max_length=100)
    issues: list[str] = Field(default_factory=list)


def document_units(content: str, filename: str) -> list[dict]:
    """保留每个非空文本行/表格单元格及来源位置，覆盖解析后的多工作表。"""
    units = []
    if filename.lower().endswith(".csv") and not any(line.strip().startswith("|") for line in content.splitlines()):
        rows = csv.reader(io.StringIO(content))
        for row_index, row in enumerate(rows, 1):
            for column, value in enumerate(row, 1):
                if value.strip():
                    units.append(
                        {"id": len(units), "location": f"csv:row:{row_index}/col:{column}", "text": value.strip()}
                    )
        return units
    for line_index, line in enumerate(content.splitlines(), 1):
        if not line.strip():
            continue
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = re.split(r"(?<!\\)\|", line.strip()[1:-1])
            if all(re.fullmatch(r"\s*:?-+:?\s*", cell) for cell in cells):
                continue
            for column, cell in enumerate(cells, 1):
                value = html.unescape(re.sub(r"<br\s*/?>", "\n", cell, flags=re.I)).replace(r"\|", "|").strip()
                if value:
                    units.append({"id": len(units), "location": f"line:{line_index}/col:{column}", "text": value})
        else:
            units.append({"id": len(units), "location": f"line:{line_index}", "text": line})
    return units


def validate_document_result(payload, document):
    result = ViralDocumentResultV1.model_validate(payload)
    if not result.articles and not result.issues:
        raise ValueError("未识别出完整文章时必须说明原因")
    units = {unit["id"]: unit for unit in document["units"]}
    used = set()
    for article in result.articles:
        if article.industry_slug not in document["industries"]:
            raise ValueError("文章行业不属于可用行业")
        if document.get("industry_hint") and article.industry_slug != document["industry_hint"]:
            raise ValueError("不得改写知识库明确配置的行业")
        ids = [article.title_unit, *article.body_units]
        if len(ids) != len(set(ids)) or any(value not in units or value in used for value in ids):
            raise ValueError("原文单元不存在、重复使用或跨文章混用")
        if article.body_units != sorted(article.body_units):
            raise ValueError("正文必须保持原文顺序")
        title_location = units[article.title_unit]["location"]
        if "/col:" in title_location:
            record = title_location.rsplit("/col:", 1)[0]
            if any(units[index]["location"].rsplit("/col:", 1)[0] != record for index in article.body_units):
                raise ValueError("表格文章的标题和正文必须来自同一记录")
        used.update(ids)
    return result


def extract_detected_articles(result, document):
    units = {unit["id"]: unit for unit in document["units"]}
    return [
        {
            "title": re.sub(r"^#{1,6}\s+", "", units[article.title_unit]["text"]),
            "body": join_body_units([units[index] for index in article.body_units]),
            "locator": units[article.title_unit]["location"],
            "industry_slug": article.industry_slug,
        }
        for article in result.articles
    ]


def join_body_units(units):
    parts = [units[0]["text"]]
    for previous, current in zip(units, units[1:]):
        before = re.fullmatch(r"line:(\d+)", previous["location"])
        after = re.fullmatch(r"line:(\d+)", current["location"])
        separator = "\n" * (int(after[1]) - int(before[1])) if before and after else "\n"
        parts.extend([separator, current["text"]])
    return "".join(parts)


def xlsx_document_units(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    units = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets, 1):
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None and str(cell.value).strip():
                        units.append(
                            {
                                "id": len(units),
                                "location": f"sheet:{sheet_index}:{sheet.title}/row:{cell.row}/col:{cell.column}",
                                "text": str(cell.value).strip(),
                            }
                        )
    finally:
        workbook.close()
    return units
